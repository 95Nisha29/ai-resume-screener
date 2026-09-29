"""Document text extraction service supporting PDF, DOCX, and TXT files."""

import io
import os
import logging
from typing import Union, BinaryIO

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_input: Union[str, bytes, BinaryIO]) -> str:
    """Extract text content from a PDF file path, raw bytes, or stream."""
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.error("pypdf is not installed.")
        raise RuntimeError("pypdf library is required for PDF parsing.")

    if isinstance(file_input, (bytes, bytearray)):
        stream = io.BytesIO(file_input)
    elif isinstance(file_input, str):
        stream = open(file_input, "rb")
    else:
        stream = file_input

    try:
        reader = PdfReader(stream)
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text.strip())
        return "\n\n".join(extracted_pages)
    except Exception as e:
        logger.warning(f"Failed to extract text from PDF: {e}")
        return ""
    finally:
        if isinstance(file_input, str) and hasattr(stream, "close"):
            stream.close()


def extract_text_from_docx(file_input: Union[str, bytes, BinaryIO]) -> str:
    """Extract text content from a DOCX file path, raw bytes, or stream."""
    try:
        import docx
    except ImportError:
        logger.error("python-docx is not installed.")
        raise RuntimeError("python-docx library is required for DOCX parsing.")

    if isinstance(file_input, (bytes, bytearray)):
        stream = io.BytesIO(file_input)
    elif isinstance(file_input, str):
        stream = file_input
    else:
        stream = file_input

    try:
        doc = docx.Document(stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        return "\n".join(paragraphs)
    except Exception as e:
        logger.warning(f"Failed to extract text from DOCX: {e}")
        return ""


def extract_text_from_txt(file_input: Union[str, bytes, BinaryIO]) -> str:
    """Extract text from plain text input with fallback encodings."""
    if isinstance(file_input, str):
        if os.path.exists(file_input):
            for encoding in ("utf-8", "latin-1", "cp1252"):
                try:
                    with open(file_input, "r", encoding=encoding, errors="replace") as f:
                        return f.read()
                except Exception:
                    continue
            return ""
        else:
            return file_input

    if isinstance(file_input, (bytes, bytearray)):
        data = file_input
    else:
        data = file_input.read()

    for encoding in ("utf-8", "latin-1", "cp1252", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue

    return data.decode("utf-8", errors="replace")


def extract_document_text(filename: str, content: Union[str, bytes, BinaryIO]) -> str:
    """Generic document text extractor based on file extension."""
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        return extract_text_from_pdf(content)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(content)
    elif ext in (".txt", ".md", ".rtf", ""):
        return extract_text_from_txt(content)
    else:
        # Fallback to plain text attempt
        return extract_text_from_txt(content)
